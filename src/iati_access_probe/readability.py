"""Readability assessment — deliverable 5 (SIGNAL-METHOD §3).

This is the deliverable-4 ContentSink's consumer. The D4 contract is
binding and is the whole reason this module looks the way it does:

  * one URL's body at a time, hard-bounded by `policy.max_bytes`;
  * assessed, its metrics written onto the SAME `CrawlResult` row;
  * the buffer released by the crawler the moment `assess_body` returns —
    nothing here keeps a reference to `body`, nothing is accumulated across
    URLs, nothing is written to disk. The committed JSONL row stays
    counts/statuses/identifiers only (location-safety / retention boundary).

D5 owns format-truth filtering. The sink hands off ANY terminal body within
the cap — including non-2xx and "declared-PDF-that's-actually-HTML" — by
design, so this module is where `detected_format` / `format_match` /
`declared_pdf_actually_html` are computed (magic-byte sniff cross-checked
against the content-type header; the header is NOT trusted) and where an
error/landing page is kept out of the document counts (`not_a_document`).

OCR — the Q3 stop-and-ask boundary, observed here:
  * Detecting that a file *needs* OCR (a scanned-image PDF with no text
    layer, or an image file) is cheap and IN scope — done inline below.
  * Actually *running* OCR to measure recoverable yield is the gate: it
    adds a dependency and possibly a per-page cost. It is NOT done here and
    no OCR library is imported. `extraction_outcome == "ocr_needed"` is the
    detected, unquantified segment; whether its yield is measured is
    SIGNAL-METHOD §4 options (b)/(c), resolved before the crawl run.

No document text is ever returned, logged, or placed on the row — only its
character count and the derived booleans (location-safety rule).
"""

from __future__ import annotations

import io

from .models import CrawlResult

# Default for Decisions Log Q6 — confirmed with Jaimie at the pre-crawl
# gate; the confirmed value is injected at run time, this is only the
# discussion default so offline tests and the sink have a value.
DEFAULT_USABLE_TEXT_THRESHOLD = 250

# How little extracted text means "this PDF has no usable text layer" — the
# trigger to look for an image (scanned page) vs. a genuinely empty doc.
# Distinct from, and far below, the usable-text threshold.
_NO_TEXT_LAYER_FLOOR = 16


# --- format detection (magic bytes + sniff, header cross-checked) -----------

# Coarse classes the declared MIME and the detected bytes are both reduced
# to; `format_match` compares at this type level (SIGNAL-METHOD §3).
_PDF, _HTML, _DOCX, _IMAGE, _TEXT, _OTHER = (
    "pdf", "html", "docx", "image", "text", "other",
)


def _class_of_mime(mime: str | None) -> str | None:
    if not mime:
        return None
    m = mime.split(";", 1)[0].strip().lower()
    if m == "application/pdf" or m.endswith("/x-pdf"):
        return _PDF
    if m in ("text/html", "application/xhtml+xml"):
        return _HTML
    if m in (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/msword",
    ):
        return _DOCX
    if m.startswith("image/"):
        return _IMAGE
    if m.startswith("text/"):
        return _TEXT
    return _OTHER


def _looks_html(head: bytes) -> bool:
    s = head[:1024].lstrip().lower()
    if s.startswith(b"<?xml"):
        s = s[s.find(b"?>") + 2 :].lstrip() if b"?>" in s else s
    return (
        s.startswith(b"<!doctype html")
        or s.startswith(b"<html")
        or b"<html" in s[:512]
    )


def _detect(body: bytes) -> tuple[str, str]:
    """Return (detected_mime, detected_class) from the bytes themselves.

    Magic-byte first (the wire header is not trusted — SIGNAL-METHOD §3);
    HTML/plain-text are sniffed because they have no binary signature.
    """
    import filetype  # pinned dep; pure-Python, no libmagic (Windows host)

    kind = filetype.guess(body)
    if kind is not None:
        cls = _class_of_mime(kind.mime)
        return kind.mime, (cls if cls is not None else _OTHER)
    if _looks_html(body):
        return "text/html", _HTML
    # Decodable, mostly-printable -> plain text; else opaque binary.
    sample = body[:4096]
    try:
        sample.decode("utf-8")
        printable = sum(c >= 0x09 for c in sample)
        if sample and printable / len(sample) > 0.85:
            return "text/plain", _TEXT
    except UnicodeDecodeError:
        pass
    return "application/octet-stream", _OTHER


# --- per-format extraction --------------------------------------------------


def _pdf_outcome(body: bytes) -> tuple[str, int]:
    """(extraction_outcome, char_count) for a PDF.

    Text layer present -> text_extracted. ~No text but a page carries an
    image -> ocr_needed (scanned). ~No text and no image -> text_extracted
    with a near-zero count (an empty/near-empty document is itself a
    finding, SIGNAL-METHOD §3). Unreadable -> extraction_failed.
    """
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(body))
        if reader.is_encrypted:
            try:
                reader.decrypt("")  # empty-password only; never a real bypass
            except Exception:
                return "extraction_failed", 0
        text_len = 0
        has_image = False
        for page in reader.pages:
            try:
                text_len += len((page.extract_text() or "").strip())
            except Exception:
                pass
            if not has_image:
                try:
                    res = page.get("/Resources")
                    xobj = res.get("/XObject") if res else None
                    if xobj and any(
                        (xobj[n].get("/Subtype") == "/Image") for n in xobj
                    ):
                        has_image = True
                except Exception:
                    pass
        if text_len >= _NO_TEXT_LAYER_FLOOR:
            return "text_extracted", text_len
        if has_image:
            return "ocr_needed", 0
        return "text_extracted", text_len  # opened, ~empty: a finding
    except Exception:
        return "extraction_failed", 0


def _docx_outcome(body: bytes) -> tuple[str, int]:
    try:
        import docx  # python-docx

        d = docx.Document(io.BytesIO(body))
        n = sum(len(p.text.strip()) for p in d.paragraphs)
        for t in d.tables:
            for row in t.rows:
                for cell in row.cells:
                    n += len(cell.text.strip())
        return "text_extracted", n
    except Exception:
        return "extraction_failed", 0


def _html_outcome(body: bytes, http_status: int | None, threshold: int) -> tuple[str, int]:
    """HTML that yields substantive text is readable (text_extracted). An
    error response or a body too thin to be a document is a nav/landing/
    error stub (not_a_document) — it must NOT be scored as a document."""
    try:
        from bs4 import BeautifulSoup

        soup = BeautifulSoup(body, "html.parser")  # stdlib parser, no lxml
        for tag in soup(["script", "style", "noscript", "template"]):
            tag.decompose()
        n = len(soup.get_text(" ", strip=True))
    except Exception:
        return "extraction_failed", 0
    if http_status is not None and http_status >= 400:
        return "not_a_document", n
    if n < threshold:
        return "not_a_document", n
    return "text_extracted", n


def _text_outcome(body: bytes) -> tuple[str, int]:
    return "text_extracted", len(body.decode("utf-8", errors="replace").strip())


# --- the assessment ---------------------------------------------------------


def assess_body(
    result: CrawlResult,
    body: bytes,
    *,
    usable_text_threshold: int = DEFAULT_USABLE_TEXT_THRESHOLD,
) -> None:
    """Fill the readability fields on `result` from `body`, in place.

    Designed to be the deliverable-4 ContentSink. `body` is this URL's
    bytes only; no reference to it is kept past this call.
    """
    detected_mime, detected_cls = _detect(body)
    declared_cls = _class_of_mime(result.declared_format)

    result.detected_format = detected_mime
    result.format_match = (
        None if declared_cls is None else (declared_cls == detected_cls)
    )
    result.declared_pdf_actually_html = (
        None
        if declared_cls is None
        else (declared_cls == _PDF and detected_cls == _HTML)
    )

    if detected_cls == _PDF:
        outcome, chars = _pdf_outcome(body)
    elif detected_cls == _DOCX:
        outcome, chars = _docx_outcome(body)
    elif detected_cls == _HTML:
        outcome, chars = _html_outcome(
            body, result.http_status, usable_text_threshold
        )
    elif detected_cls == _IMAGE:
        outcome, chars = "ocr_needed", 0  # image, no text layer
    elif detected_cls == _TEXT:
        outcome, chars = _text_outcome(body)
    else:
        outcome, chars = "extraction_failed", 0

    result.extraction_outcome = outcome
    result.extracted_char_count = chars
    result.usable_text = (
        outcome == "text_extracted" and chars >= usable_text_threshold
    )
    result.reachable_and_readable = bool(
        result.resolved
        and result.access_barrier == "none"
        and (result.usable_text or outcome == "ocr_needed")
    )
    result.readability_pending = False


def mark_no_body(result: CrawlResult) -> None:
    """Finalise a row the sink never reached (robots_disallowed,
    non_public_host, oversize, any error, a redirect-only dead end). No body
    was assessed, so it cannot be reachable-and-readable."""
    result.readability_pending = False
    result.extraction_outcome = None
    result.extracted_char_count = None
    result.usable_text = False
    result.reachable_and_readable = False


def readability_sink(usable_text_threshold: int = DEFAULT_USABLE_TEXT_THRESHOLD):
    """Build the on_content callable the crawler expects (ContentSink)."""

    def _sink(result: CrawlResult, body: bytes) -> None:
        assess_body(result, body, usable_text_threshold=usable_text_threshold)

    return _sink
